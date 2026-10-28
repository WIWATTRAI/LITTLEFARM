# -*- coding: utf-8 -*-
"""ตารางโครงสร้างรายจุด: ฐานราก เสา คาน (ขนาด ระดับ ความลึก ปริมาณ) + ช่องกรอกเหล็กเสริมตามเอกสารออกแบบ
   out/02_ตารางโครงสร้าง_ฐานราก_เสา_คาน.xlsx
ใช้งาน: python3 make_struct_schedule.py
- ข้อมูลเรขาคณิตทุกจุดมาจาก model.py (ตำแหน่งจากเวกเตอร์ S201/S202 ขนาดจากตาราง S1.1/S3/S4.1 แบบ A103/A104)
- เหล็กเสริม: ช่องสีเหลืองในชีต "*-เหล็กเสริม" กรอกตามเอกสารออกแบบ แล้วน้ำหนักทุกจุดคำนวณด้วยสูตร Excel
  น้ำหนักเหล็ก (กก./ม.) = Ø² / 162 (Ø เป็นมม.)
"""
import os, math, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
import model as M

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out", "02_ตารางโครงสร้าง_ฐานราก_เสา_คาน.xlsx")
E = M.build()
S = M.schedule()
F = "Tahoma"
thin = Side(style="thin", color="999999"); BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="1F3A4D"); YEL = PatternFill("solid", fgColor="FFF2CC"); SUB = PatternFill("solid", fgColor="E8EEF0")
PEND = PatternFill("solid", fgColor="E7E6F7")
LEAN, SAND, WS = 0.05, 0.10, 0.10          # คอนกรีตหยาบ, ทรายรองพื้น, ระยะเผื่อข้างงานขุด (สมมติ)
TITLE = "งานก่อสร้างอาคารโกดังและงานบริเวณ บริษัท ตงเฉิน อีเล็คทริค เคเบิ้ล แอนด์ไวร์ จำกัด | WEC-TC-6901"

piles = [e for e in E if e["cat"].startswith("pile")]
PILE_NAME = {"I35-22": "I-350 ยาว 22 ม.", "I26-22": "I-260 ยาว 22 ม.", "I26": "I-260 (ถัง)", "I22-22": "I-220 ยาว 22 ม.", "I22-TR": "I-220 ยาว 22 ม.", "I26-B": "I-260 ยาว 22 ม."}


def piles_under(e):
    cx, cy = e["c"][0], e["c"][1]; ha, hb = e["a"]/2 + 0.05, e["b"]/2 + 0.05
    if abs(e["ex"][0]) < 0.5: ha, hb = hb, ha
    got = [p for p in piles if abs(p["c"][0]-cx) <= ha and abs(p["c"][1]-cy) <= hb]
    c = collections.Counter(p["mark"] for p in got)
    return sum(c.values()), ", ".join(f"{n}x {PILE_NAME.get(k, k)}" for k, n in c.items())


def top(e): return e["c"][2] + e["h"]/2
def bot(e): return e["c"][2] - e["h"]/2


wb = Workbook()


def header(ws, title, cols, widths, row=3):
    ws.append([title]); ws["A1"].font = Font(name=F, bold=True, size=13)
    ws.append([TITLE])
    ws.append(cols)
    for i, c in enumerate(cols, 1):
        cl = ws.cell(row=row, column=i); cl.fill = HDR; cl.font = Font(name=F, bold=True, color="FFFFFF", size=9)
        cl.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for i, w in enumerate(widths, 1): ws.column_dimensions[CL(i)].width = w
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = f"C{row+1}"


def finish(ws, first=4):
    for r in ws.iter_rows(min_row=first, max_row=ws.max_row):
        for cl in r:
            if cl.value is None and cl.fill.fgColor.rgb in (None, "00000000"): pass
            cl.border = BRD
            if not cl.font.bold: cl.font = Font(name=F, size=9, bold=cl.font.bold)
            if isinstance(cl.value, float): cl.number_format = "#,##0.000"


# ============================================================ FOOTINGS
FT_TYPES = [  # mark, name, cover
    ("F2P", "ฐานราก F2P 0.60x1.50x0.80 (เข็ม 2 ต้น) รับเสา C1/C2", 0.075),
    ("F1P", "Drop panel F1P 1.00x1.00x0.20 ใต้พื้น S1 (เข็ม I35 1 ต้น)", 0.050),
    ("F1P-1", "ฐาน F1P-1 0.60x0.60x0.60 อาคารประกอบ (เข็ม 1 ต้น)", 0.075),
    ("F1A", "ฐาน F1A 0.60x0.60x0.30 ใต้พื้นถนน S2 (เข็ม I22 1 ต้น)", 0.050),
    ("SQ66", "ฐาน 0.66x0.66 แนวเขตที่ดิน (หน้าที่รอยืนยัน)", 0.050),
    ("TRP", "แท่นหม้อแปลง 750 kVA (A104) ฐาน 5.00x3.00x0.40 + แท่น 3.50x1.75x0.15", 0.075),
]
ws = wb.active; ws.title = "1-ฐานราก รายจุด"
cols = ["รหัส", "ชนิด", "ตำแหน่งกริด", "พิกัด x (ม.)", "พิกัด y (ม.)", "กว้าง X (ม.)", "ยาว Y (ม.)", "หนา (ม.)", "ระดับหลังฐาน", "ระดับท้องฐาน",
        "ระดับก้นหลุม (หัก คสล.หยาบ+ทราย)", "ความลึกขุดจาก ±0.00 (ม.)", "เข็ม (ต้น)", "ชนิดเข็ม", "คอนกรีต (ลบ.ม.)", "ไม้แบบ (ตร.ม.)",
        "ดินขุด (ลบ.ม.)", "ทรายรองพื้น (ลบ.ม.)", "คอนกรีตหยาบ (ลบ.ม.)", "เหล็กเสริม (กก.)", "กิจกรรม 4D", "สถานะ"]
header(ws, "ตารางฐานรากรายจุด: ขนาด ระดับ ความลึก เข็ม และปริมาณ (เหล็กเสริมคำนวณจากชีต 2)", cols,
       [9, 8, 13, 9, 9, 8, 8, 7, 9, 9, 11, 10, 6, 18, 10, 9, 9, 9, 9, 11, 8, 26])
ft_rows = collections.defaultdict(list)
for mk, nm, cov in FT_TYPES:
    els = sorted([e for e in E if e.get("id", "").startswith(mk + "-") and e["mark"] == mk], key=lambda e: e["id"])
    for e in els:
        if mk == "TRP" and e["ifc"] != "IfcFooting": continue
        a, b = (e["a"], e["b"]) if abs(e["ex"][0]) > 0.5 else (e["b"], e["a"])
        t, bt = top(e), bot(e); pit = bt - LEAN - SAND
        dep = max(0.0, -pit); n, ptxt = piles_under(e)
        r = ws.max_row + 1
        exc = f"=IF(L{r}>0,(F{r}+2*{WS})*(G{r}+2*{WS})*L{r},0)"
        status = e["src"]
        ws.append([e["id"], mk, e["grid"], round(e["c"][0], 3), round(e["c"][1], 3), round(a, 3), round(b, 3), round(e["h"], 3), round(t, 3), round(bt, 3),
                   round(pit, 3), round(dep, 3), n, ptxt, f"=F{r}*G{r}*H{r}", f"=2*(F{r}+G{r})*H{r}", exc,
                   f"=(F{r}+0.2)*(G{r}+0.2)*{SAND}", f"=(F{r}+0.2)*(G{r}+0.2)*{LEAN}",
                   f"=INDEX('2-ฐานราก เหล็กเสริม'!$S:$S,MATCH(B{r},'2-ฐานราก เหล็กเสริม'!$A:$A,0))", e["task"], status])
        ft_rows[mk].append(r)
        if "รอยืนยัน" in status or "สมมติ" in status: ws.cell(row=r, column=22).fill = PEND
last_ft = ws.max_row
ws.append([]); ws.append(["รวม", "", "", "", "", "", "", "", "", "", "", "", f"=SUM(M4:M{last_ft})", "", f"=SUM(O4:O{last_ft})", f"=SUM(P4:P{last_ft})",
                          f"=SUM(Q4:Q{last_ft})", f"=SUM(R4:R{last_ft})", f"=SUM(S4:S{last_ft})", f"=SUM(T4:T{last_ft})"])
for c in ws[ws.max_row]: c.font = Font(name=F, bold=True, size=9)
ws.append(["หมายเหตุ: ระดับท้องฐานตามที่ใช้ในโมเดล (F2P -1.20 จากความยาวเสา C1 4.90 ม. ตาราง S4.1, F1P +0.10 ชิดท้องพื้น S1, F1A -0.45 ชิดท้องพื้น S2, F1P-1 -1.00, แท่นหม้อแปลง -0.40 ตาม A104) "
           f"ยังรอยืนยันจากแบบขยายฐานราก; ดินขุดคิดจากระดับดินเดิม ±0.00 เผื่อข้าง {WS} ม. ใต้ฐานมีคอนกรีตหยาบ {LEAN} + ทราย {SAND} ม."])
finish(ws)

# ---------------------------------------------------------- footing rebar inputs
w2 = wb.create_sheet("2-ฐานราก เหล็กเสริม")
c2 = ["ชนิด", "รายละเอียด", "ด้านยาว (ม.)", "ด้านสั้น (ม.)", "หนา (ม.)", "ระยะหุ้ม (ม.)",
      "ตะแกรงล่าง แนวยาว: จำนวน", "Ø (มม.)", "ตะแกรงล่าง แนวสั้น: จำนวน", "Ø (มม.)",
      "ตะแกรงบน แนวยาว: จำนวน", "Ø (มม.)", "ตะแกรงบน แนวสั้น: จำนวน", "Ø (มม.)",
      "เหล็กรัดรอบข้าง: จำนวนชั้น", "Ø (มม.)", "งอปลาย (เท่าของ Ø)", "จำนวนฐาน", "เหล็กเสริม/ฐาน (กก.)", "เหล็กเสริมรวม (กก.)", "สถานะ"]
header(w2, "เหล็กเสริมฐานราก แยกตามชนิด (กรอกช่องสีเหลืองตามเอกสารออกแบบ ระบบคำนวณน้ำหนักทุกฐานอัตโนมัติ)", c2,
       [8, 40, 8, 8, 7, 7, 9, 7, 9, 7, 9, 7, 9, 7, 9, 7, 8, 8, 11, 12, 22])
for mk, nm, cov in FT_TYPES:
    e = next(x for x in E if x["mark"] == mk and x.get("id") and x["ifc"] == "IfcFooting")
    L_, W_ = max(e["a"], e["b"]), min(e["a"], e["b"])
    r = w2.max_row + 1
    # น้ำหนัก/ฐาน: เหล็กแนวยาว ยาว = L-2c+2*k*Ø ; แนวสั้น = W-2c+2*k*Ø ; รัดรอบ = 2(L+W)-8c+0.3 ต่อชั้น
    kg = (f"=G{r}*(C{r}-2*F{r}+2*Q{r}*H{r}/1000)*H{r}^2/162 + I{r}*(D{r}-2*F{r}+2*Q{r}*J{r}/1000)*J{r}^2/162"
          f" + K{r}*(C{r}-2*F{r}+2*Q{r}*L{r}/1000)*L{r}^2/162 + M{r}*(D{r}-2*F{r}+2*Q{r}*N{r}/1000)*N{r}^2/162"
          f" + O{r}*(2*(C{r}+D{r})-8*F{r}+0.3)*P{r}^2/162")
    w2.append([mk, nm, round(L_, 2), round(W_, 2), round(e["h"], 2), cov, None, None, None, None, None, None, None, None, None, None, 12,
               f"=COUNTIF('1-ฐานราก รายจุด'!$B:$B,A{r})", kg, f"=R{r}*S{r}", "รอเอกสารออกแบบเหล็กเสริม"])
    for c in range(7, 18): w2.cell(row=r, column=c).fill = YEL
    w2.cell(row=r, column=21).fill = PEND
r = w2.max_row + 1
w2.append(["รวม"] + [None]*17 + [None, f"=SUM(T4:T{r-1})"])
w2.append([]); w2.append(["ตัวอย่างการกรอก: F2P ตะแกรงล่างแนวยาว 6 เส้น Ø16 → ช่อง G=6, H=16; ตะแกรงล่างแนวสั้น 10 เส้น Ø12 → I=10, J=12 (ระบุจำนวนเส้นจริงตามแบบ ไม่ใช่ระยะห่าง)"])
w2.append(["เหล็กยืนของเสาที่ฝังในฐานราก คิดในชีต 4 (เสา) ไม่ซ้ำในชีตนี้"])
finish(w2)

# ============================================================ COLUMNS
w3 = wb.create_sheet("3-เสา รายต้น")
c3 = ["รหัส", "ชนิด", "ตำแหน่งกริด", "พิกัด x", "พิกัด y", "หน้าตัด X (ม.)", "หน้าตัด Y (ม.)", "ระดับโคนเสา", "ระดับหัวเสา", "ความสูง (ม.)",
      "ฐานรากรองรับ", "คอนกรีต (ลบ.ม.)", "ไม้แบบ (ตร.ม.)", "เหล็กเสริม (กก.)", "สถานะ"]
header(w3, "ตารางเสา คสล. รายต้น (โคนเสา = หลังฐานราก, หัวเสา = +4.50 / อาคารประกอบ +3.00)", c3, [9, 6, 12, 8, 8, 9, 9, 9, 9, 9, 12, 10, 10, 11, 28])
foots = [e for e in E if e["ifc"] == "IfcFooting" and e.get("id")]
for mk in ("C1", "C2", "C3", "SC"):
    for e in sorted([x for x in E if x["mark"] == mk and x.get("id")], key=lambda x: x["id"]):
        f_ = next((f for f in foots if abs(f["c"][0]-e["c"][0]) < 0.3 and abs(f["c"][1]-e["c"][1]) < 0.3), None)
        r = w3.max_row + 1
        w3.append([e["id"], mk, e["grid"], round(e["c"][0], 3), round(e["c"][1], 3), round(e["a"], 3), round(e["b"], 3), round(bot(e), 3), round(top(e), 3),
                   f"=I{r}-H{r}", f_["id"] if f_ else ("ผนัง/คานคอดิน" if mk == "SC" else "-"), f"=F{r}*G{r}*J{r}",
                   (f"=2*(F{r}+G{r})*J{r}" if mk != "SC" else f"=2*MIN(F{r},G{r})*J{r}"),
                   f"=INDEX('4-เสา เหล็กเสริม'!$O:$O,MATCH(B{r},'4-เสา เหล็กเสริม'!$A:$A,0))", e["src"]])
last_c = w3.max_row
w3.append(["รวม"] + [None]*10 + [f"=SUM(L4:L{last_c})", f"=SUM(M4:M{last_c})", f"=SUM(N4:N{last_c})"])
finish(w3)

w4 = wb.create_sheet("4-เสา เหล็กเสริม")
c4 = ["ชนิด", "รายละเอียด", "หน้าตัด b (ม.)", "หน้าตัด h (ม.)", "ระยะหุ้ม (ม.)", "เหล็กยืน: จำนวน", "Ø (มม.)", "ระยะทาบ/ฝัง (เท่าของ Ø)",
      "งอฝังในฐาน (ม.)", "ปลอก Ø (มม.)", "ระยะปลอก ช่วงปลาย (ม.)", "ระยะปลอก ช่วงกลาง (ม.)", "ความสูงเฉลี่ย (ม.)", "จำนวนต้น", "เหล็กเสริม/ต้น (กก.)", "เหล็กเสริมรวม (กก.)", "สถานะ"]
header(w4, "เหล็กเสริมเสา แยกตามชนิด (กรอกช่องสีเหลือง)", c4, [7, 40, 8, 8, 7, 9, 7, 9, 9, 8, 9, 9, 9, 8, 11, 12, 22])
for mk, nm in [("C1", "เสา C1 0.30x0.50 จากหลังฐาน F2P ถึง +4.50"), ("C2", "เสา C2 0.50x0.50 (6/J)"), ("C3", "เสา C3 0.20x0.20 อาคารประกอบ"),
               ("SC", "เสาเอ็น 0.15x0.10 ในผนังก่อ +0.50 ถึง +4.35")]:
    e = next(x for x in E if x["mark"] == mk and x.get("id"))
    b_, h_ = sorted([e["a"], e["b"]])
    r = w4.max_row + 1
    # ยืน: n*(H + la*Ø + งอฝัง) ; ปลอก: จำนวน = 2*ROUNDUP(H/4/s1)+ROUNDUP(H/2/s2)+1 ; ยาว = 2(b+h)-8c+24Øs
    kg = (f"=F{r}*(M{r}+H{r}*G{r}/1000+I{r})*G{r}^2/162"
          f" + IF(J{r}=0,0,(2*ROUNDUP(M{r}/4/K{r},0)+ROUNDUP(M{r}/2/L{r},0)+1)*(2*(C{r}+D{r})-8*E{r}+24*J{r}/1000)*J{r}^2/162)")
    w4.append([mk, nm, round(b_, 2), round(h_, 2), 0.04 if mk != "SC" else 0.02, None, None, 40, 0.30 if mk in ("C1", "C2", "C3") else 0.0, None, None, None,
               f"=AVERAGEIF('3-เสา รายต้น'!$B:$B,A{r},'3-เสา รายต้น'!$J:$J)", f"=COUNTIF('3-เสา รายต้น'!$B:$B,A{r})", kg, f"=N{r}*O{r}", "รอเอกสารออกแบบเหล็กเสริม"])
    for c in (6, 7, 8, 9, 10, 11, 12): w4.cell(row=r, column=c).fill = YEL
    w4.cell(row=r, column=17).fill = PEND
r = w4.max_row + 1
w4.append(["รวม"] + [None]*14 + [f"=SUM(P4:P{r-1})"])
w4.append([]); w4.append(["หมายเหตุ: ถ้าไม่กรอกระยะปลอก ระบบไม่คิดปลอก (ป้องกันหารด้วยศูนย์) ช่วงปลาย = 1/4 ของความสูงแต่ละด้าน"])
finish(w4)

# ============================================================ BEAMS
w5 = wb.create_sheet("5-คาน รายช่วง")
c5 = ["รหัส", "ชนิด", "แนว", "จาก (กริด)", "ถึง (กริด)", "ยาว ศก.-ศก. (ม.)", "ยาวสุทธิ ผิว-ผิว (ม.)", "กว้าง b (ม.)", "ลึก h (ม.)", "ระดับหลังคาน", "ระดับท้องคาน",
      "คอนกรีต (ลบ.ม.)", "ไม้แบบ (ตร.ม.)", "เหล็กบน (กก.)", "เหล็กล่าง (กก.)", "เหล็กเสริมพิเศษ (กก.)", "ปลอก (กก.)", "ปลอก (ตัว)", "เหล็กรวม (กก.)", "สถานะ"]
header(w5, "ตารางคาน คสล. รายช่วง: หน้าตัด ความยาว ระดับ ปริมาณ และเหล็กเสริม (สูตรอ้างชีต 6)", c5,
       [9, 6, 10, 11, 11, 9, 9, 7, 7, 9, 9, 9, 9, 9, 9, 10, 9, 8, 10, 26])
REF = "'6-คาน เหล็กเสริม'!"
def idx(col, r): return f"INDEX({REF}${col}:${col},MATCH($B{r},{REF}$A:$A,0))"
SUPW = {"B1": 0.30, "B3": 0.20, "LB": 0.30, "TB": 0.30}
for mk in ("B1", "B3", "LB", "TB"):
    for e in sorted([x for x in E if x["mark"] == mk and x.get("id")], key=lambda x: x["id"]):
        p0 = M.sub(e["c"], M.mul(e["ez"], e["h"]/2)); p1 = M.add(e["c"], M.mul(e["ez"], e["h"]/2))
        along = "X" if abs(e["ez"][0]) > 0.5 else "Y"
        line = (f"แนว {M.grid_tag(0, p0[1]).split('/')[1]}" if along == "X" else f"แนว {M.grid_tag(p0[0], 0).split('/')[0]}")
        Lseg = e["h"]
        if mk in ("LB", "TB"): Lcc, Lcl = Lseg + 0.30, Lseg
        else: Lcc, Lcl = Lseg, Lseg - SUPW[mk]
        b_, h_ = e["a"], e["b"]
        r = w5.max_row + 1
        la = f"{idx('Q', r)}"
        top_ = f"={idx('E',r)}*(F{r}+2*{idx('Q',r)}*{idx('F',r)}/1000)*{idx('F',r)}^2/162"
        botm = f"={idx('G',r)}*(F{r}+2*{idx('Q',r)}*{idx('H',r)}/1000)*{idx('H',r)}^2/162"
        extra = (f"={idx('I',r)}*2*(0.3*G{r}+{idx('Q',r)}*{idx('J',r)}/1000)*{idx('J',r)}^2/162"
                 f"+{idx('K',r)}*0.7*G{r}*{idx('L',r)}^2/162")
        nst = f"=IF({idx('M',r)}=0,0,2*ROUNDUP(G{r}/4/{idx('N',r)},0)+ROUNDUP(G{r}/2/{idx('O',r)},0)+1)"
        stw = f"=R{r}*(2*(H{r}+I{r})-8*{idx('P',r)}+24*{idx('M',r)}/1000)*{idx('M',r)}^2/162"
        w5.append([e["id"], mk, line, M.grid_tag(p0[0], p0[1]), M.grid_tag(p1[0], p1[1]), round(Lcc, 3), round(Lcl, 3), round(b_, 2), round(h_, 2),
                   round(e["c"][2] + h_/2, 3), round(e["c"][2] - h_/2, 3), f"=H{r}*I{r}*G{r}",
                   (f"=(2*I{r}+H{r})*G{r}" if mk in ("B1", "B3") else f"=2*I{r}*G{r}"), top_, botm, extra, stw, nst, f"=SUM(N{r}:Q{r})", e["src"]])
last_b = w5.max_row
w5.append(["รวม"] + [None]*4 + [f"=SUM(F4:F{last_b})", f"=SUM(G4:G{last_b})"] + [None]*4 + [f"=SUM(L4:L{last_b})", f"=SUM(M4:M{last_b})",
          f"=SUM(N4:N{last_b})", f"=SUM(O4:O{last_b})", f"=SUM(P4:P{last_b})", f"=SUM(Q4:Q{last_b})", f"=SUM(R4:R{last_b})", f"=SUM(S4:S{last_b})"])
w5.append(["หมายเหตุ: B1 ยาวสุทธิ = ศก.-ศก. หัก 0.30 ม. (ด้านแคบของเสา C1); B3 หัก 0.20 ม. (C3); คานทับหลัง LB / คานเอ็น TB วัดผิวเสาถึงผิวเสา; "
           "ไม้แบบ B1/B3 = 2 ข้าง + ท้อง, LB/TB = 2 ข้าง (วางบนผนัง)"])
finish(w5)

w6 = wb.create_sheet("6-คาน เหล็กเสริม")
c6 = ["ชนิด", "รายละเอียด", "b (ม.)", "h (ม.)", "เหล็กบนต่อเนื่อง: จำนวน", "Ø (มม.)", "เหล็กล่างต่อเนื่อง: จำนวน", "Ø (มม.)",
      "เสริมบนที่หัวเสา: จำนวน", "Ø (มม.)", "เสริมล่างกลางช่วง: จำนวน", "Ø (มม.)", "ปลอก Ø (มม.)", "ระยะปลอก ช่วงริม L/4 (ม.)", "ระยะปลอก ช่วงกลาง (ม.)",
      "ระยะหุ้ม (ม.)", "ระยะทาบ/ฝัง (เท่าของ Ø)", "จำนวนช่วง", "ยาวรวม (ม.)", "เหล็กรวม (กก.)", "สถานะ"]
header(w6, "เหล็กเสริมคาน แยกตามชนิด: เหล็กยืน เหล็กเสริมพิเศษ และเหล็กปลอก (กรอกช่องสีเหลือง)", c6,
       [6, 42, 6, 6, 9, 6, 9, 6, 9, 6, 9, 6, 7, 9, 9, 7, 9, 7, 9, 11, 22])
for mk, nm in [("B1", "คานคอดิน B1 0.20x0.60 หลังคาน +0.50 (รอบอาคาร + แนวกริด 6)"), ("B3", "คาน B3 0.20x0.40 อาคารประกอบ (ห้องน้ำ +0.50 / ห้องสารเคมี +0.40)"),
               ("LB", "คานทับหลัง 0.15x0.15 +4.35..+4.50 บนผนังก่อ"), ("TB", "คานเอ็น 0.15x0.10 ระดับหลังวงกบ +2.60")]:
    e = next(x for x in E if x["mark"] == mk and x.get("id"))
    r = w6.max_row + 1
    w6.append([mk, nm, round(e["a"], 2), round(e["b"], 2)] + [None]*11 + [0.04 if mk in ("B1", "B3") else 0.02, 40,
               f"=COUNTIF('5-คาน รายช่วง'!$B:$B,A{r})", f"=SUMIF('5-คาน รายช่วง'!$B:$B,A{r},'5-คาน รายช่วง'!$F:$F)",
               f"=SUMIF('5-คาน รายช่วง'!$B:$B,A{r},'5-คาน รายช่วง'!$S:$S)", "รอเอกสารออกแบบเหล็กเสริม"])
    for c in range(5, 18): w6.cell(row=r, column=c).fill = YEL
    w6.cell(row=r, column=21).fill = PEND
r = w6.max_row + 1
w6.append(["รวม"] + [None]*18 + [f"=SUM(T4:T{r-1})"])
w6.append([]); w6.append(["วิธีคิด: เหล็กยืนต่อเนื่อง ยาว = ศก.-ศก. + 2 x ระยะทาบ/ฝัง x Ø; เสริมบนที่หัวเสา 2 ปลาย/ช่วง ยาวปลายละ 0.3 x ยาวสุทธิ + ระยะฝัง; "
          "เสริมล่าง 0.7 x ยาวสุทธิ; ปลอก ช่วงริม L/4 สองด้าน + ช่วงกลาง L/2, ยาวต่อตัว = 2(b+h) - 8 x ระยะหุ้ม + ขอ 24Ø"])
finish(w6)

# ============================================================ STEEL COLUMN BASE (D01)
w8 = wb.create_sheet("8-รอยต่อเสาเหล็ก")
c8 = ["รหัสเสาเหล็ก", "ตำแหน่งกริด", "เสา คสล. ใต้", "ด้านยาวแผ่นตามแนว", "Base plate (มม.)", "Base plate (กก.)", "Top plate (มม.)", "Top plate (กก.)",
      "Anchor bolt (ชุด)", "Anchor bolt (กก.)", "ระยะฝังสลัก (ม.)", "เกราท์ (ลบ.ม.)", "ระยะหุ้มเหล็กยืนเสา (ม.)", "ระยะสลัก-เหล็กมุมเสา ศก. (มม.)", "ระยะช่องว่าง (มม.)", "ผล"]
header(w8, "รอยต่อเสาเหล็ก H244x175 กับหัวเสา คสล.: แผ่นเหล็กฐาน PL20x200x450 + Anchor bolt M20x750 6 ชุด + แผ่นหัวเสา PL20x200x450 (แบบขยาย D01)", c8,
       [10, 10, 9, 9, 13, 9, 13, 9, 8, 9, 9, 9, 10, 11, 10, 18])
C1R = "'4-เสา เหล็กเสริม'!"
rcs = [e for e in E if e["mark"] in ("C1", "C2") and e.get("id")]
bolt_u, bolt_v = max(abs(u) for u, v in M.BP["bolts"]), max(abs(v) for u, v in M.BP["bolts"])
for e in sorted([x for x in E if x["mark"] == "SC1" and x.get("id")], key=lambda x: x["id"]):
    rc = min(rcs, key=lambda r_: abs(r_["c"][0]-e["c"][0]) + abs(r_["c"][1]-e["c"][1]))
    long_x = rc["a"] > rc["b"] or (rc["a"] == rc["b"] and rc["mark"] == "C2" and False)
    hb, hh = (min(rc["a"], rc["b"]), max(rc["a"], rc["b"]))
    r = w8.max_row + 1
    typ = rc["mark"]
    # ระยะจากศูนย์กลางถึงเหล็กมุม = ครึ่งด้าน - ระยะหุ้ม - ปลอก - Ø/2 (อ่านจากชีต 4 ตามชนิดเสา)
    cov = f"INDEX({C1R}$E:$E,MATCH(\"{typ}\",{C1R}$A:$A,0))"; dia = f"INDEX({C1R}$G:$G,MATCH(\"{typ}\",{C1R}$A:$A,0))"
    tie = f"INDEX({C1R}$J:$J,MATCH(\"{typ}\",{C1R}$A:$A,0))"
    du = f"({hh/2}-{cov}-{tie}/1000-{dia}/2000-{bolt_u})"; dv = f"({hb/2}-{cov}-{tie}/1000-{dia}/2000-{bolt_v})"
    w8.append([e["id"], e["grid"], rc["id"], "X" if rc["a"] >= rc["b"] else "Y", "20x200x450", M.BP_KG, "20x200x450", M.BP_KG,
               len(M.BP["bolts"]), round(len(M.BP["bolts"])*M.AB_KG, 2), round(M.BP["bolt_L"]-M.BP["bolt_proj"]-M.BP["t"]-M.BP["grout"], 3),
               round((M.BP["l"]+0.05)*(M.BP["w"]+0.05)*M.BP["grout"], 4), f"={cov}",
               f"=IF({dia}=0,\"รอตารางเสา\",ROUND(1000*SQRT({du}^2+{dv}^2),1))",
               f"=IF({dia}=0,\"รอตารางเสา\",ROUND(N{r}-{M.BP['bolt_d']*1000/2}-{dia}/2,1))",
               f"=IF({dia}=0,\"รอตารางเสา\",IF(O{r}>=25,\"ผ่าน (≥25 มม.)\",\"ชน/แคบ ตรวจแบบ\"))"])
last8 = w8.max_row
w8.append(["รวม", "", "", "", f"=COUNTA(E4:E{last8})", f"=SUM(F4:F{last8})", f"=COUNTA(G4:G{last8})", f"=SUM(H4:H{last8})", f"=SUM(I4:I{last8})", f"=SUM(J4:J{last8})", "", f"=SUM(L4:L{last8})"])
w8.append(["หมายเหตุ: ค่าเบื้องต้นจากผู้ว่าจ้าง 28/10/69; ตำแหน่งสลัก 3x2 ระยะ 175x100 มม. ระยะขอบแผ่น 50 มม. (สมมติ) ฝังจากหัวเสา +4.50 ลึก 645 มม.; "
           "การตรวจระยะใช้เหล็กยืนมุมเสา (ระยะหุ้ม + ปลอก + Ø/2) จากชีต 4 เกณฑ์ช่องว่าง ≥ 25 มม.; สลักคู่กลางอยู่ระหว่างปีกเสา H ตรวจที่ขันน็อต"])
finish(w8)

# ============================================================ CHECK / OPEN ISSUES
w7 = wb.create_sheet("7-ข้อมูลที่ต้องยืนยัน")
w7.append(["ข้อมูลที่ต้องยืนยันหรือขอจากผู้ออกแบบก่อนขึ้นรูปเหล็กเสริมและสรุปราคา"]); w7["A1"].font = Font(name=F, bold=True, size=13)
w7.append(["ลำดับ", "ชิ้นส่วน", "ประเด็น", "ข้อมูลในแบบปัจจุบัน", "สิ่งที่ต้องการ", "ผลต่อ BOQ"])
for c in range(1, 7):
    cl = w7.cell(row=2, column=c); cl.fill = HDR; cl.font = Font(name=F, bold=True, color="FFFFFF", size=9)
ISSUES = [
    ("ฐานราก F2P", "ระดับหลังฐาน/ท้องฐาน และระยะเข็มในฐาน", "ตาราง S1.1 ขนาด 0.60x1.50x0.80; โมเดลใช้ท้องฐาน -1.20 จากความยาวเสา C1", "แบบขยายฐานราก (รูปตัด ระดับ ระยะเข็ม ระยะหุ้ม)", "ดินขุด ไม้แบบ ค้ำยัน"),
    ("ฐานรากทุกชนิด", "เหล็กเสริมฐานราก", "ไม่มีในแบบ", "ตะแกรงล่าง/บน จำนวน-ขนาด ทั้ง 2 แนว เหล็กรัดรอบ (กรอกชีต 2)", "2.2.4 เหล็กเสริมฐานราก"),
    ("Drop panel F1P", "ความสัมพันธ์กับพื้น S1 (หล่อพร้อมพื้นหรือแยก) และเหล็กเสริม", "1.00x1.00x0.20 ใต้ท้องพื้น S1 +0.30", "แบบขยาย Drop panel + เหล็กเสริม", "2.2.2 / 2.4.3"),
    ("เสา C1/C2/C3", "เหล็กยืนและปลอก", "ไม่มีในแบบ", "จำนวน-ขนาดเหล็กยืน ปลอก Ø@ระยะ (กรอกชีต 4)", "2.3.6"),
    ("คาน B1/B3", "เหล็กบน ล่าง เสริมพิเศษ และปลอก", "หน้าตัดจากตาราง S3 เท่านั้น", "หน้าตัดยืนยัน + เหล็กเสริมแต่ละช่วง (กรอกชีต 6)", "2.3.6"),
    ("คานทับหลัง/คานเอ็น/เสาเอ็น", "หน้าตัดและเหล็กเสริมมาตรฐาน", "ผู้ว่าจ้างกำหนดระดับ +4.35/+4.50; หน้าตัดสมมติ", "รายละเอียดมาตรฐานผนังก่อ", "3.2.3-3.2.6"),
    ("รอยต่อเสาเหล็ก-หัวเสา (D01)", "ระยะ/ตำแหน่งสลัก การต่อแผ่นหัวเสากับโครงถัก", "ผู้ว่าจ้างกำหนด PL20x200x450 + M20x750 6 ชุด; ตำแหน่งสลักสมมติ 175x100 มม.", "ยืนยันระยะสลัก ระยะขอบ ขนาดรู และสลักยึดโครงถัก", "2.5.7-2.5.9"),
    ("แท่นหม้อแปลง (A104)", "จำนวนและตำแหน่งเข็มไม่ตรงกัน", "A104 แสดงเข็ม 3x3 = 9 ต้น ระยะ 2.10 x 1.10 ม.; S201/S201-B แสดง I22-TR 6 ต้น ระยะ ~1.50 ม.", "ยืนยันจำนวน/ตำแหน่งเข็ม และระดับแท่น", "4.3.1 (+3 ต้นถ้าใช้ A104)"),
    ("ห้องเก็บสารเคมี (A103)", "ระดับพื้นและตำแหน่งเสา", "A103 พื้น +0.40 มีเสากลางผนังที่ y=3.0 ม.; S201 มี F1P-1 ที่ y=0, 6.0, 9.0 (ไม่มีที่ y=3.0); S6.1 AF01 ระดับถนน", "ยืนยันตำแหน่งฐาน/เสา และระดับพื้น", "2.2.3 / 2.3.3 / 2.4.2"),
    ("เข็มถนน I22", "ตำแหน่งที่ขาด 17 ต้น", "ตาราง S2 = 249, ผัง S201/S202 = 232", "ตำแหน่งเข็ม 17 ต้น", "4.1.1 / 4.1.2 (ใช้ 249 แล้ว)"),
    ("ฐาน 0.66 ม. แนวเขต 16 จุด", "หน้าที่และระดับ", "พบในผัง S201 ไม่มีในตาราง", "ระบุว่าเป็นฐานอะไร (รั้ว/เสาไฟ/ป้าย) และขนาดความลึก", "4.6.1"),
    ("ถังเก็บน้ำใต้ดิน", "ผนัง ฝา เหล็กเสริม", "S202 แสดงเฉพาะพื้น TF -2.60 / BF -2.85", "แบบขยายถัง", "4.4.3 / 4.4.4"),
]
for i, row in enumerate(ISSUES, 1): w7.append([i] + list(row))
for i, w in enumerate([6, 22, 30, 50, 44, 22], 1): w7.column_dimensions[CL(i)].width = w
for rr in w7.iter_rows(min_row=3):
    for cl in rr: cl.border = BRD; cl.font = Font(name=F, size=9); cl.alignment = Alignment(wrap_text=True, vertical="top")

# ============================================================ SUMMARY (first sheet)
sm = wb.create_sheet("สรุป", 0)
sm.append(["สรุปตารางโครงสร้าง ฐานราก เสา คาน รายจุด"]); sm["A1"].font = Font(name=F, bold=True, size=13)
sm.append([TITLE])
sm.append(["กลุ่ม", "ชนิด", "จำนวน", "คอนกรีต (ลบ.ม.)", "ไม้แบบ (ตร.ม.)", "เหล็กเสริม (กก.)", "เทียบรายการ BOQ v2"])
for c in range(1, 8):
    cl = sm.cell(row=3, column=c); cl.fill = HDR; cl.font = Font(name=F, bold=True, color="FFFFFF", size=9)
BOQ = {"F2P": "2.2.1", "F1P": "2.2.2", "F1P-1": "2.2.3", "F1A": "4.1.2", "SQ66": "4.6.1", "TRP": "4.3.2",
       "C1": "2.3.1", "C2": "2.3.2", "C3": "2.3.3", "SC": "3.2.5", "B1": "2.3.4", "B3": "2.3.5", "LB": "3.2.3", "TB": "3.2.4"}
for mk, *_ in FT_TYPES:
    sm.append(["ฐานราก", mk, f"=COUNTIF('1-ฐานราก รายจุด'!B:B,B{sm.max_row+1})", f"=SUMIF('1-ฐานราก รายจุด'!B:B,B{sm.max_row+1},'1-ฐานราก รายจุด'!O:O)",
               f"=SUMIF('1-ฐานราก รายจุด'!B:B,B{sm.max_row+1},'1-ฐานราก รายจุด'!P:P)", f"=SUMIF('1-ฐานราก รายจุด'!B:B,B{sm.max_row+1},'1-ฐานราก รายจุด'!T:T)", BOQ[mk]])
for mk in ("C1", "C2", "C3", "SC"):
    sm.append(["เสา", mk, f"=COUNTIF('3-เสา รายต้น'!B:B,B{sm.max_row+1})", f"=SUMIF('3-เสา รายต้น'!B:B,B{sm.max_row+1},'3-เสา รายต้น'!L:L)",
               f"=SUMIF('3-เสา รายต้น'!B:B,B{sm.max_row+1},'3-เสา รายต้น'!M:M)", f"=SUMIF('3-เสา รายต้น'!B:B,B{sm.max_row+1},'3-เสา รายต้น'!N:N)", BOQ[mk]])
for mk in ("B1", "B3", "LB", "TB"):
    sm.append(["คาน", mk, f"=COUNTIF('5-คาน รายช่วง'!B:B,B{sm.max_row+1})", f"=SUMIF('5-คาน รายช่วง'!B:B,B{sm.max_row+1},'5-คาน รายช่วง'!L:L)",
               f"=SUMIF('5-คาน รายช่วง'!B:B,B{sm.max_row+1},'5-คาน รายช่วง'!M:M)", f"=SUMIF('5-คาน รายช่วง'!B:B,B{sm.max_row+1},'5-คาน รายช่วง'!S:S)", BOQ[mk]])
n = sm.max_row
sm.append(["รวม", "", f"=SUM(C4:C{n})", f"=SUM(D4:D{n})", f"=SUM(E4:E{n})", f"=SUM(F4:F{n})", ""])
sm.append([]); sm.append(["วิธีใช้: กรอกเหล็กเสริมในชีต 2 (ฐานราก), 4 (เสา), 6 (คาน) ตามเอกสารออกแบบ แล้วน้ำหนักเหล็กทุกจุดในชีต 1, 3, 5 และสรุปนี้คำนวณใหม่ทันที; "
                         "รหัสชิ้นส่วนตรงกับโมเดล 3D (คลิกชิ้นส่วนเพื่อดูรหัส) และ Pset_TC_BOQ_4D ใน IFC"])
for i, w in enumerate([10, 8, 9, 13, 12, 14, 18], 1): sm.column_dimensions[CL(i)].width = w
for rr in sm.iter_rows(min_row=4, max_row=n+1):
    for cl in rr:
        cl.border = BRD; cl.font = Font(name=F, size=10, bold=cl.row == n+1)
        if cl.column in (4, 5, 6): cl.number_format = "#,##0.00"
wb.save(OUT)
print(OUT)
print({k: len(v) for k, v in ft_rows.items()}, "beams", last_b-3, "columns", last_c-3)
